import csv
import json
import os
import subprocess
import sys
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from book_manager.services.auditoria import auditar_operaciones


@auditar_operaciones
class ServicioScraping:
    def __init__(self, conexion):
        self._conexion = conexion
        self.raiz = Path(__file__).resolve().parents[3]
        self.proyecto_scrapy = (
            self.raiz / "src" / "book_manager" / "scraper"
            / "cuspide_scrapy"
        )
        self.carpeta_reportes = (
            self.raiz / "src" / "book_manager" / "reportes"
        )
        self.carpeta_reportes.mkdir(parents=True, exist_ok=True)

    def ejecutar_scraping(self, umbral_ars):
        try:
            umbral = Decimal(str(umbral_ars))
        except (InvalidOperation, ValueError) as error:
            raise ValueError("Ingresá un umbral numérico en ARS.") from error
        if not umbral.is_finite() or umbral < 0:
            raise ValueError("El umbral debe ser cero o mayor.")

        marca = datetime.now().strftime("%Y%m%d_%H%M%S")
        ruta_json = self.carpeta_reportes / f"cuspide_datos_{marca}.json"
        ruta_csv = self.carpeta_reportes / f"alertas_precios_{marca}.csv"
        entorno = os.environ.copy()
        fuente = str(self.raiz / "src")
        entorno["PYTHONPATH"] = fuente + os.pathsep + entorno.get("PYTHONPATH", "")

        resultado = subprocess.run(
            [sys.executable, "-m", "scrapy", "crawl", "cuspide",
             "-O", str(ruta_json)],
            cwd=self.proyecto_scrapy,
            env=entorno,
            capture_output=True,
            text=True,
            timeout=3600,
            check=False,
        )
        if resultado.returncode != 0:
            salida = (resultado.stderr or resultado.stdout)[-2500:]
            raise OSError(f"Falló CuspideSpider:\\n{salida}")
        if not ruta_json.exists():
            raise OSError("Scrapy no generó el JSON de resultados.")

        filas = json.loads(ruta_json.read_text(encoding="utf-8"))
        alertas = []
        for fila in filas:
            interno = Decimal(str(fila["precio_interno"]))
            web = Decimal(str(fila["precio_web"]))
            diferencia = web - interno
            if abs(diferencia) >= umbral:
                alertas.append({
                    "libro": fila["titulo_catalogo"],
                    "precio_interno_ars": f"{interno:.2f}",
                    "precio_web_ars": f"{web:.2f}",
                    "diferencia_web_menos_interno_ars": f"{diferencia:.2f}",
                    "umbral_ars": f"{umbral:.2f}",
                    "fecha_extraccion": fila["fecha_extraccion"],
                    "url_producto": fila["url_producto"],
                })

        campos = [
            "libro", "precio_interno_ars", "precio_web_ars",
            "diferencia_web_menos_interno_ars", "umbral_ars",
            "fecha_extraccion", "url_producto",
        ]
        with ruta_csv.open("w", newline="", encoding="utf-8-sig") as archivo:
            escritor = csv.DictWriter(archivo, fieldnames=campos)
            escritor.writeheader()
            escritor.writerows(alertas)
        return {
            "libros": len(filas), "alertas": len(alertas),
            "csv": ruta_csv, "json": ruta_json,
        }

    def generar_reporte_excel(self):
        import json
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        archivos = sorted(
            self.carpeta_reportes.glob("cuspide_datos_*.json"),
            key=lambda ruta: ruta.stat().st_mtime,
        )
        if not archivos:
            raise FileNotFoundError(
                "Primero ejecutá el scraping desde el menú."
            )
        datos = json.loads(archivos[-1].read_text(encoding="utf-8"))
        libro_excel = Workbook()
        hoja = libro_excel.active
        hoja.title = "Precios Cúspide"
        hoja.append([
            "Libro", "Precio interno", "Precio web",
            "Diferencia", "Fecha de extracción",
        ])
        for celda in hoja[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="234D3D")

        for fila in datos:
            interno = Decimal(str(fila["precio_interno"]))
            web = Decimal(str(fila["precio_web"]))
            hoja.append([
                fila["titulo_catalogo"], float(interno), float(web),
                float(web - interno), fila["fecha_extraccion"],
            ])
        for fila in hoja.iter_rows(min_row=2, min_col=2, max_col=4):
            for celda in fila:
                celda.number_format = '#,##0.00'
        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = hoja.dimensions
        for columna in hoja.columns:
            ancho = max(len(str(c.value or "")) for c in columna)
            letra = get_column_letter(columna[0].column)
            hoja.column_dimensions[letra].width = min(ancho + 2, 55)

        marca = datetime.now().strftime("%Y%m%d_%H%M%S")
        destino = self.carpeta_reportes / f"reporte_cuspide_{marca}.xlsx"
        libro_excel.save(destino)
        return destino